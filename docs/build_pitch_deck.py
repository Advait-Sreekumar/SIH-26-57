"""Build docs/SonarEye_Pitch_Deck.pptx from the approved 11-slide outline.

Dark theme to match the Part B diagrams (bg #1a1a2e).
Run from repo root:  python docs/build_pitch_deck.py

All numbers here trace to source files; no mine-detection language;
the synthetic-net slide carries the mandatory caveat label verbatim.
"""
from __future__ import annotations
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

ROOT = Path(__file__).resolve().parent.parent
DIAGRAMS = Path(__file__).resolve().parent / "diagrams"
OUT = Path(__file__).resolve().parent / "SonarEye_Pitch_Deck.pptx"

# palette (matches generate_diagrams.py)
DARK_BG = RGBColor(0x1A, 0x1A, 0x2E)
ACCENT  = RGBColor(0xE9, 0x45, 0x60)
LIGHT   = RGBColor(0xE0, 0xE0, 0xE0)
GREY    = RGBColor(0x9A, 0x9A, 0xB0)
AMBER   = RGBColor(0xFF, 0x98, 0x00)
TEAL    = RGBColor(0x26, 0xA6, 0x9A)

# 16:9
SW, SH = Inches(13.333), Inches(7.5)

prs = Presentation()
prs.slide_width = SW
prs.slide_height = SH
BLANK = prs.slide_layouts[6]

DIAG = {
    "01": DIAGRAMS / "01_system_architecture.png",
    "02": DIAGRAMS / "02_confidence_score_breakdown.png",
    "03": DIAGRAMS / "03_evaluation_charts.png",
    "04": DIAGRAMS / "04_class_expansion_decisions.png",
}


def _bg(slide):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = DARK_BG


def _title(slide, text, y=Inches(0.35), color=LIGHT, size=30):
    tb = slide.shapes.add_textbox(Inches(0.6), y, Inches(12.1), Inches(1.0))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = True
    r.font.color.rgb = color
    r.font.name = "Segoe UI"
    return tb


def _bullets(slide, items, x=Inches(0.7), y=Inches(1.6), w=Inches(7.4),
             h=Inches(5.2), size=17):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    first = True
    for item in items:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        # item may be (text, color) or plain str
        if isinstance(item, tuple):
            text, col = item
        else:
            text, col = item, LIGHT
        r = p.add_run()
        r.text = "•  " + text
        r.font.size = Pt(size)
        r.font.color.rgb = col
        r.font.name = "Segoe UI"
        p.space_after = Pt(10)
        p.line_spacing = 1.05
    return tb


def _pic_fit(slide, img_path, x, y, max_w, max_h):
    """Add picture scaled to fit inside (max_w, max_h), centered in that box."""
    from PIL import Image
    with Image.open(img_path) as im:
        iw, ih = im.size
    ar = iw / ih
    box_ar = max_w / max_h
    if ar > box_ar:
        w = max_w
        h = Emu(int(max_w * ih / iw))
    else:
        h = max_h
        w = Emu(int(max_h * iw / ih))
    px = Emu(int(x + (max_w - w) / 2))
    py = Emu(int(y + (max_h - h) / 2))
    slide.shapes.add_picture(str(img_path), px, py, width=w, height=h)


def _caption(slide, text, y, color=GREY, size=11, x=Inches(0.6),
             w=Inches(12.1), italic=True, align=PP_ALIGN.CENTER):
    tb = slide.shapes.add_textbox(x, y, w, Inches(0.5))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.italic = italic
    r.font.color.rgb = color
    r.font.name = "Segoe UI"
    return tb


def _table(slide, rows, x, y, w, col_widths, header=True):
    """Simple dark-themed table from a list of row-lists (strings)."""
    n_rows = len(rows)
    n_cols = len(rows[0])
    h = Inches(0.42 * n_rows)
    gtbl = slide.shapes.add_table(n_rows, n_cols, x, y, w, h).table
    for ci, cw in enumerate(col_widths):
        gtbl.columns[ci].width = cw
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = gtbl.cell(ri, ci)
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor(0x0F, 0x34, 0x60) if (header and ri == 0) else RGBColor(0x22, 0x22, 0x3A)
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            r = p.add_run()
            r.text = val
            r.font.size = Pt(12)
            r.font.color.rgb = LIGHT
            r.font.bold = (header and ri == 0)
            r.font.name = "Segoe UI"
    return gtbl


def build():
    # Slide 1 — Title / Problem
    s = prs.slides.add_slide(BLANK); _bg(s)
    tb = s.shapes.add_textbox(Inches(0.8), Inches(2.4), Inches(11.7), Inches(1.4))
    p = tb.text_frame.paragraphs[0]
    r = p.add_run(); r.text = "SonarEye"
    r.font.size = Pt(54); r.font.bold = True; r.font.color.rgb = ACCENT; r.font.name = "Segoe UI"
    sub = s.shapes.add_textbox(Inches(0.8), Inches(3.7), Inches(11.7), Inches(0.8))
    pr = sub.text_frame.paragraphs[0]
    rr = pr.add_run(); rr.text = "Finding what doesn't belong on the seafloor"
    rr.font.size = Pt(24); rr.font.color.rgb = LIGHT; rr.font.name = "Segoe UI"
    _bullets(s, [
        "Ghost nets and marine debris keep catching life long after they're abandoned — but on side-scan sonar they blend into natural clutter (rock, reef, sand ripples).",
        "Sonar survey review is still largely manual: an analyst scrubs hours of imagery frame by frame.",
        "The hard problem: telling a real anomalous object from high-contrast natural texture, at scale, without missing things.",
    ], y=Inches(4.6), w=Inches(11.8), size=16)

    # Slide 2 — What we built and proved
    s = prs.slides.add_slide(BLANK); _bg(s)
    _title(s, "What we built and proved")
    _bullets(s, [
        "End-to-end pipeline: raw sonar → preprocessing → tiled U-Net segmentation → post-processing → geotagging → mission report, with a human review loop.",
        "Validated on real shipwreck side-scan sonar (AI4Shipwrecks: 286 images, 28 sites; site-based split, zero site overlap between train and test).",
        ("Headline: Test IoU 0.4268 / Dice 0.5983 on 13 held-out sites.", ACCENT),
        "In-distribution validation reaches IoU 0.713 / Dice 0.833 — the gap is reported, not hidden.",
    ], y=Inches(1.5), w=Inches(6.4), size=16)
    _pic_fit(s, DIAG["01"], Inches(7.3), Inches(1.6), Inches(5.7), Inches(4.9))

    # Slide 3 — System architecture
    s = prs.slides.add_slide(BLANK); _bg(s)
    _title(s, "How it works, end to end")
    _pic_fit(s, DIAG["01"], Inches(0.5), Inches(1.35), Inches(12.3), Inches(4.6))
    _caption(s, "Preprocessing (Lee filter + water crop + CLAHE) · U-Net ResNet34, 512×512 tiles @50% overlap · "
                "connected-components + sonar-heuristic confidence · pixel→lat/lon geotagging · JSON/CSV/PDF report · "
                "SQLite-persisted human review loop", Inches(6.15), size=12)

    # Slide 4 — Evidence: evaluation & robustness
    s = prs.slides.add_slide(BLANK); _bg(s)
    _title(s, "Evidence: evaluation & robustness")
    _pic_fit(s, DIAG["03"], Inches(0.5), Inches(1.3), Inches(12.3), Inches(4.1))
    _bullets(s, [
        ("The single convergent finding: the model fires on high-contrast structured texture, "
         "not exclusively on shipwreck shape.", AMBER),
        "The perturbations that degrade it most (speckle −0.0453, heave/pitch/roll −0.0254) are the ones that alter "
        "texture; four independent stress tests point the same way. Baseline IoU 0.0581.",
    ], y=Inches(5.5), w=Inches(12.3), size=14)

    # Slide 5 — Evidence: deployment & performance
    s = prs.slides.add_slide(BLANK); _bg(s)
    _title(s, "Evidence: deployment & performance")
    _table(s, [
        ["Metric", "Measured value", "Scope"],
        ["GPU forward pass", "0.37 s", "pure model, warm"],
        ["Full pipeline (GPU)", "686 ms", "preprocess + infer + post, warm"],
        ["Full pipeline (CPU)", "4.29 s", "warm"],
        ["Cold start", "~7,120 ms", "incl. ~5.9 s one-time imports"],
        ["ONNX FP32 export", "97.7 MB · max diff 2.57e-05", "PASS"],
        ["ONNX INT8 export", "24.6 MB · diff 0.6497", "WARN — not used"],
    ], Inches(0.7), Inches(1.6), Inches(11.9),
       [Inches(3.2), Inches(4.9), Inches(3.8)])
    _bullets(s, [
        "Runs today as a Streamlit app.",
        ('No "real-time" or "edge-ready" claim — those require measurement on target hardware, which we have not done.', AMBER),
    ], y=Inches(5.4), w=Inches(11.9), size=15)

    # Slide 6 — Taxonomy gap
    s = prs.slides.add_slide(BLANK); _bg(s)
    _title(s, "The taxonomy gap: investigated, not hidden")
    _pic_fit(s, DIAG["04"], Inches(0.5), Inches(1.3), Inches(12.3), Inches(4.5))
    _caption(s, "Shipwrecks SHIPPED (test IoU 0.427) · Pipes+cylinders BLOCKED (SubPipe GPL-3.0 + domain shift) · "
                "Ghost nets BLOCKED (no public labeled sonar data) · Generic debris DEFERRED. "
                "The system boundary is defined by evidence, not omission.", Inches(6.05), size=12)

    # Slide 7 — Class-agnostic architecture
    s = prs.slides.add_slide(BLANK); _bg(s)
    _title(s, "The pipeline doesn't know what it's looking at")
    _bullets(s, [
        "The segmentation model outputs \"anomalous object vs seafloor\" — it is not hardcoded to \"shipwreck.\"",
        "Adding a second class is a data problem, not an architecture rewrite: the review workflow already stores a "
        "generic category field (shipwreck, rock, shadow_artifact, biological, unknown).",
        ("Shipwrecks were simply the class that had real, licensed, labeled data to prove the architecture on. "
         "The design generalizes; the evidence is class-specific by necessity.", TEAL),
    ], y=Inches(1.7), w=Inches(12.0), size=18)

    # Slide 8 — Synthetic net exploration (caveated)
    s = prs.slides.add_slide(BLANK); _bg(s)
    _title(s, "Synthetic net exploration — a texture probe")
    _bullets(s, [
        "Exploratory probe on synthetic net-like texture: 26/30 (86.7%) patches triggered a detection "
        "(72 detections total, seed 7).",
        "Consistent with the Slide 4 finding — the model fires on unfamiliar high-contrast texture, not on net shape.",
    ], y=Inches(1.7), w=Inches(12.0), size=18)
    # mandatory caveat box
    box = s.shapes.add_textbox(Inches(1.0), Inches(4.6), Inches(11.3), Inches(1.4))
    box.fill.solid() if False else None
    tf = box.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = ("Label: “illustrative, unvalidated, synthetic-only exploration — not a demonstrated capability.”  "
              "This is NOT ghost-net detection and is not presented as one.")
    r.font.size = Pt(16); r.font.bold = True; r.font.italic = True
    r.font.color.rgb = AMBER; r.font.name = "Segoe UI"

    # Slide 9 — Human-in-the-loop workflow
    s = prs.slides.add_slide(BLANK); _bg(s)
    _title(s, "AI assists, a human verifies")
    _bullets(s, [
        "Mission flow: landing → survey → analysis → review-in-sequence → map → summary → report.",
        "Each detection is reviewed in sequence: confirm / reject / uncertain, with optional category and note; "
        "decisions persist across restarts (SQLite).",
        ("The report distinguishes AI-flagged / unreviewed from human-confirmed — the tool never presents an "
         "unreviewed detection as a verified finding.", TEAL),
    ], y=Inches(1.7), w=Inches(12.0), size=18)

    # Slide 10 — What real deployment would need next
    s = prs.slides.add_slide(BLANK); _bg(s)
    _title(s, "What real deployment would need next")
    _bullets(s, [
        "Labeled ghost-net sonar data — the one missing ingredient. Plausible route: partner with a conservation / "
        "ghost-gear-retrieval organization to co-label real survey imagery, the way AI4Shipwrecks itself was built.",
        "Embedded-hardware benchmarking before any \"edge\" or \"real-time\" claim can be made.",
        "Validation on Indian-waters sonar to confirm generalization beyond freshwater Great Lakes data.",
    ], y=Inches(1.7), w=Inches(12.0), size=18)

    # Slide 11 — Close
    s = prs.slides.add_slide(BLANK); _bg(s)
    _title(s, "What's real today, and what's next")
    _bullets(s, [
        ("Real and tested today: end-to-end pipeline validated on real held-out shipwreck sonar (Test IoU 0.4268), "
         "with measured latency, a human review loop, and geotagged reports.", LIGHT),
        ("Honestly bounded: we documented every class we couldn't ship and why; the net work is an unvalidated "
         "texture probe, not a capability.", LIGHT),
        ("Next: labeled ghost-net data + hardware benchmarking + regional validation turn a proven architecture "
         "into a deployable tool.", ACCENT),
    ], y=Inches(1.8), w=Inches(12.0), size=19)

    prs.save(str(OUT))
    print(f"Saved: {OUT}  ({len(prs.slides)} slides)")


if __name__ == "__main__":
    build()




